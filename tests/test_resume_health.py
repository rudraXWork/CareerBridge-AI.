from io import BytesIO

import pymupdf
import pytest
from docx import Document
from docx.oxml.ns import qn
from fastapi.testclient import TestClient

from backend.main import app
from backend.models import CandidateProfile, JobInput
from backend.nlp.job_parser import parse_job
from backend.nlp.matcher import match_resume
from backend.nlp.resume_health import analyze_resume, inspect_document
from backend.nlp.resume_parser import extract_text, parse_resume


def check(report, key):
    return next(item for item in report['ats']['checks'] if item['id'] == key)


def test_health_requires_no_job_or_providers(candidate, monkeypatch):
    monkeypatch.setattr('backend.services.web.request_json', lambda *args, **kwargs: pytest.fail('Health is local'))
    res = TestClient(app).post('/resume/analyze', json={'candidate': candidate.model_dump()})
    assert res.status_code == 200
    report = res.json()['health']
    assert 0 <= report['score'] <= 100
    assert report['ats']['coverage'] < 100
    assert check(report, 'tables')['status'] == 'unknown'
    assert 'job' not in report


def test_empty_profile_does_not_get_a_positive_score():
    report = analyze_resume(CandidateProfile())
    assert report['score'] == 0
    assert report['ats']['score'] == 0
    assert report['content']['score'] == 0


def test_empty_section_headings_do_not_get_full_ats_credit():
    text = 'Asha Rao\nasha@example.com\nSKILLS\nEDUCATION\nPROJECTS'
    profile = parse_resume(text)
    profile.document = inspect_document(text.encode(), 'resume.txt', text)
    assert profile.document.section_names == []
    report = analyze_resume(profile)
    assert check(report, 'sections')['points'] == 0
    assert report['ats']['score'] == 80
    # Rechecking an older upload must also ignore its empty section headings.
    profile.document.section_names = ['skills', 'education', 'projects']
    assert check(analyze_resume(profile), 'sections')['points'] == 0


def test_pdf_100_means_all_assessed_checks_passed():
    text = ('Asha Rao\nasha@example.com\nSKILLS\nPython SQL\nEDUCATION\nBachelor of Computer Science\n'
            'PROJECTS\nBuilt a dashboard using Python and SQL to track sales.')
    doc = pymupdf.open()
    doc.new_page().insert_text((40, 60), text)
    content = doc.tobytes()
    doc.close()
    extracted = extract_text(content, 'resume.pdf')
    profile = parse_resume(extracted)
    profile.document = inspect_document(content, 'resume.pdf', extracted)
    ats = analyze_resume(profile)['ats']
    assert ats['score'] == 100
    assert ats['coverage'] == 75
    assert sum(item['points'] or 0 for item in ats['checks']) == 75
    assert {item['id'] for item in ats['checks'] if item['status'] == 'unknown'} == {
        'tables', 'text_boxes', 'header_footer_contact'}
    profile.email = ''
    assert analyze_resume(profile)['ats']['score'] == 90


def test_student_projects_count_and_edits_ignore_stale_sections(candidate):
    candidate.projects = ['Dashboard: Worked on a dashboard.']
    before = analyze_resume(candidate)
    candidate.projects = ['Dashboard: Built a Power BI dashboard using SQL to report on 50,000 sales records.']
    after = analyze_resume(candidate)
    assert after['content']['score'] > before['content']['score']
    assert after['ats']['score'] == before['ats']['score']
    assert any(item['id'] == 'metrics' and item['status'] == 'pass' for item in after['content']['checks'])
    assert not candidate.experience
    assert after['content']['checks'][-1]['status'] == 'pass'


def test_dates_versions_and_project_titles_are_not_impact(candidate):
    candidate.projects = ['Dashboard', 'Built a dashboard using Python 3.11 in 2023.']
    report = analyze_resume(candidate)
    metrics = next(item for item in report['content']['checks'] if item['id'] == 'metrics')
    assert metrics['points'] == 0
    assert '0 of 1' in metrics['evidence']


def test_health_uses_complete_bullets_not_project_metadata(candidate):
    candidate.projects = [
        'AIQuant | React, Node.js, Python, MongoDB GitHub ↗ LIVE ↗',
        '• Delivered a market intelligence application with',
        'real-time dashboards.',
        '• Built six regression models with an R2 score of 96.06%',
        'on Vercel and Render.',
        'Agent Control Plane | Python, FastAPI, SQLite, Next.js, React',
        '• Developed a governance gateway for technical policy checks.',
        'MoneyTrail AI',
        '• Analyzed transaction data using Python for financial crime detection.',
    ]
    checks = {item['id']: item for item in analyze_resume(candidate)['content']['checks']}
    assert checks['actions']['evidence'] == '4 of 4 statements start with a recognized action verb.'
    assert checks['metrics']['evidence'] == '1 of 4 statements contain a measurable result or scale.'
    assert checks['metrics']['example'] == 'Delivered a market intelligence application with real-time dashboards.'
    assert 'Example:' not in checks['metrics']['evidence']
    assert checks['actions']['example'] == ''


def test_health_keeps_colon_descriptions_and_bullet_colons(candidate):
    candidate.projects = ['Dashboard: Built a dashboard using SQL to track sales.',
                          '• Improved response time: cached 500 requests using Redis.']
    checks = {item['id']: item for item in analyze_resume(candidate)['content']['checks']}
    assert checks['actions']['evidence'].startswith('2 of 2')
    assert checks['metrics']['example'] == 'Built a dashboard using SQL to track sales.'


def test_captured_and_engineered_are_action_verbs(candidate):
    candidate.projects = ['• Captured audit events using Python for usage reports.',
                          '• Engineered behavioral features using SQL for anomaly detection.']
    actions = next(item for item in analyze_resume(candidate)['content']['checks'] if item['id'] == 'actions')
    assert actions['status'] == 'pass'
    assert actions['example'] == ''


def test_employer_dates_and_awards_are_not_content_statements():
    text = ('Asha Rao\nasha@example.com\nSKILLS\nPython SQL\n'
            'EXPERIENCE\nIndia Space Lab Jun 2026 – Jul 2026\n'
            'Summer Internship & Technical Training Program Remote\n'
            '• Trained in drone technology and remote sensing.\n'
            'Industry Capstone Project – Demo Nov 2025 – Jan 2026\n'
            'AI Decision Support & Explainability\n'
            '• Built a decision support system using Python\nand model explanations.\n'
            'Selected Projects\nDemo Real-Time Surveillance & Identity Recognition\n'
            '• Developed a surveillance pipeline using Python.\n'
            'Achievements & Certifications\n• 2nd Place – Campus Hackathon 2026.\n')
    profile = parse_resume(text)
    assert profile.projects[0].startswith('Demo Real-Time')
    assert 'achievements' in profile.sections
    assert not any('2nd Place' in line for line in profile.experience)
    content = analyze_resume(profile)['content']['checks']
    actions = next(item for item in content if item['id'] == 'actions')
    assert actions['evidence'].startswith('3 of 3')
    assert all('India Space Lab' not in item['example'] for item in content)


def test_pdf_bullet_glyph_on_its_own_line(candidate):
    candidate.projects = ['•', 'Built a dashboard using SQL', 'to track sales.', '•']
    actions = next(item for item in analyze_resume(candidate)['content']['checks'] if item['id'] == 'actions')
    assert actions['evidence'].startswith('1 of 1')


def test_duplicate_bullets_have_evidence(candidate):
    bullet = 'Built a dashboard using SQL to analyze sales trends.'
    candidate.projects = [bullet, bullet]
    repeated = next(item for item in analyze_resume(candidate)['content']['checks'] if item['id'] == 'repetition')
    assert repeated['status'] == 'review'
    assert bullet in repeated['evidence']


def test_docx_layout_is_inspected_and_survives_profile_edits():
    doc = Document()
    doc.add_paragraph('Asha Rao')
    doc.add_paragraph('SKILLS\nPython SQL')
    doc.add_paragraph('EDUCATION\nBachelor of Computer Science')
    doc.add_paragraph('PROJECTS\nBuilt a dashboard using Python and SQL.')
    doc.add_table(rows=1, cols=1).cell(0, 0).text = 'Portfolio link'
    doc.sections[0].header.paragraphs[0].text = 'asha@example.com'
    doc.sections[0]._sectPr.xpath('./w:cols')[0].set(qn('w:num'), '2')
    buf = BytesIO()
    doc.save(buf)
    text = extract_text(buf.getvalue(), 'resume.docx')
    profile = parse_resume(text)
    profile.document = inspect_document(buf.getvalue(), 'resume.docx', text)
    report = analyze_resume(profile)
    assert check(report, 'tables')['status'] == 'review'
    assert check(report, 'columns')['status'] == 'review'
    assert check(report, 'header_footer_contact')['status'] == 'review'
    assert report['ats']['coverage'] == 100
    profile.email = 'asha@example.com'
    updated = analyze_resume(profile)
    assert updated['ats']['score'] > report['ats']['score']
    assert check(updated, 'header_footer_contact') == check(report, 'header_footer_contact')


def test_pdf_column_warning_is_a_heuristic_and_tables_unknown():
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_textbox((40, 60, 250, 170), 'Python SQL skills and projects\nBuilt a useful sales dashboard')
    page.insert_textbox((310, 60, 530, 170), 'Education and experience details\nBachelor of Computer Science')
    content = doc.tobytes()
    doc.close()
    text = extract_text(content, 'resume.pdf')
    profile = parse_resume(text)
    profile.document = inspect_document(content, 'resume.pdf', text)
    report = analyze_resume(profile)
    assert check(report, 'columns')['status'] == 'review'
    assert 'Possible' in check(report, 'columns')['evidence']
    assert check(report, 'tables')['status'] == 'unknown'
    assert report['ats']['coverage'] < 100


def test_upload_includes_health_and_can_recheck(candidate):
    client = TestClient(app)
    text = ('Asha Rao\nasha@example.com\nSKILLS\nPython SQL\nEDUCATION\nBachelor of Computer Science\n'
            'PROJECTS\nBuilt a dashboard using Python and SQL to track sales.')
    res = client.post('/resume/upload', files={'file': ('resume.txt', text.encode(), 'text/plain')})
    assert res.status_code == 200
    data = res.json()
    assert data['profile']['document']['format'] == 'TXT'
    assert data['health']['ats']['coverage'] == 100
    assert client.post('/resume/analyze', json={'candidate': data['profile']}).json()['health'] == data['health']
    assert client.post('/resume/analyze', json={}).status_code == 422


def test_health_and_job_match_are_independent(candidate, jobs, monkeypatch):
    monkeypatch.setattr('backend.nlp.matcher.similarities', lambda texts, target: [0.8] * len(texts))
    health = analyze_resume(candidate)
    first = match_resume(candidate, parse_job(JobInput(**jobs[0])))
    second = match_resume(candidate, parse_job(JobInput(**jobs[2])))
    assert first.score != second.score
    assert analyze_resume(candidate) == health
