import json
from io import BytesIO

import pymupdf
import pytest
from docx import Document

from backend.config import ROOT
from backend.models import JobInput
from backend.nlp.job_parser import parse_job
from backend.nlp.matcher import education_score, experience_score, match_resume
from backend.nlp.resume_parser import clean_text, detect_sections, extract_text, parse_resume
from backend.nlp.skills import extract_skills, normalize_skills

labels = json.loads((ROOT / 'examples' / 'skill_labels.json').read_text())


@pytest.mark.parametrize('case', labels)
def test_labelled_skills(case):
    assert set(extract_skills(case['text'])[1]) == set(case['skills'])


def test_normalization_and_boundaries():
    assert normalize_skills(['JS', 'javascript', 'Postgres', 'Sklearn']) == ['JavaScript', 'PostgreSQL', 'Scikit-learn']
    assert extract_skills('JavaScript PostgreSQL GitHub Pythonic pythons go to work')[1] == ['JavaScript', 'PostgreSQL']


def test_cleaning_and_sections():
    text = clean_text('Asha Rao\r\n\r\nSKILLS:\r\n C++   C#\xa0.NET\n\n\nProjects: Node.js app')
    sections = detect_sections(text)
    assert sections['skills'] == 'C++ C# .NET'
    assert sections['projects'] == 'Node.js app'
    assert clean_text(text) == text


def test_resume_profile_and_project_skills(candidate):
    assert candidate.name == 'Asha Rao'
    assert candidate.email == 'asha@example.com'
    assert 'PostgreSQL' in candidate.normalized_skills
    assert 'Data Analyst' in candidate.possible_roles
    assert candidate.experience_years is None
    assert candidate.projects and candidate.education


def test_missing_headings_and_explicit_experience():
    profile = parse_resume('Priya Kumar\nPython and SQL\nBuilt a dashboard with Power BI')
    assert set(profile.normalized_skills) == {'Python', 'SQL', 'Power BI'}
    profile = parse_resume('Priya Kumar\nSUMMARY\n3 years of professional experience\nEXPERIENCE\nSQL data analyst')
    assert profile.experience_years == 3


def test_full_header_name_over_partial_ner():
    text = ('Rudra Narayan Jena\n'
            '(+91) 9000000000 | rudra@example.com | linkedin.com/in/rudra-jena | github.com/rudraXWork\n'
            'SKILLS\nJavaScript React')
    assert parse_resume(text).name == 'Rudra Narayan Jena'


def test_pdf_docx_and_txt_extraction():
    doc = pymupdf.open()
    doc.new_page().insert_text((50, 50), 'Asha Rao Python SQL')
    assert 'Python SQL' in extract_text(doc.tobytes(), 'resume.pdf')
    doc.close()
    word = Document()
    word.add_paragraph('Asha Rao')
    table = word.add_table(rows=1, cols=1)
    table.cell(0, 0).text = 'Python SQL'
    data = BytesIO()
    word.save(data)
    assert 'Python SQL' in extract_text(data.getvalue(), 'resume.docx')
    assert extract_text(b'\xef\xbb\xbfPython SQL', 'resume.txt') == 'Python SQL'


@pytest.mark.parametrize('content,filename', [(b'', 'a.txt'), (b'x', 'a.exe'), (b'bad pdf', 'a.pdf'),
                                             (b'bad docx', 'a.docx'), (b'\xff', 'a.txt')])
def test_unreadable_uploads(content, filename):
    with pytest.raises(ValueError):
        extract_text(content, filename)


def test_job_parsing(jobs):
    job = parse_job(JobInput(**jobs[0]))
    assert set(job.required_skills) == {'Python', 'SQL', 'Pandas', 'Power BI'}
    assert set(job.preferred_skills) == {'Excel', 'Git'}
    assert job.responsibilities
    assert job.seniority == 'Entry level'
    senior = parse_job(JobInput(**jobs[2]))
    assert senior.min_experience_years == 5
    assert senior.seniority == 'Senior'
    optional = parse_job(JobInput(job_description='Requirements\nPython required.\nPreferred\n3 years of work experience preferred.'))
    assert optional.min_experience_years is None
    assert parse_job(JobInput(job_description='Preferred\nMaster degree preferred.')).education_required == ''


def test_weighted_score_and_profile_corrections(candidate, jobs, monkeypatch):
    monkeypatch.setattr('backend.nlp.matcher.similarities', lambda texts, target: [0.8] * len(texts))
    job = parse_job(JobInput(**jobs[0]))
    result = match_resume(candidate, job)
    assert result.components['skills'] == 100
    assert result.components['experience'] == 100
    assert result.components['education'] == 100
    expected = sum(result.components[key] * weight for key, weight in result.weights_used.items())
    assert result.score == round(expected, 2)
    assert result.strong_matches and not result.missing_skills
    candidate.normalized_skills = ['SQL']
    edited = match_resume(candidate, job)
    assert edited.components['skills'] < result.components['skills']
    assert 'Python' in edited.missing_skills


def test_unknown_components_and_missing_model(candidate, jobs, monkeypatch):
    monkeypatch.setattr('backend.nlp.matcher.similarities', lambda texts, target: [0.7] * len(texts))
    job = parse_job(JobInput(company='Demo', role='Analyst', job_description=jobs[0]['job_description'].split('Education')[0]))
    result = match_resume(candidate, job)
    assert result.components['education'] is None
    assert abs(sum(result.weights_used.values()) - 1) < 1e-9
    assert result.evidence_coverage < 100

    def unavailable(*args):
        raise OSError('MiniLM unavailable')

    monkeypatch.setattr('backend.nlp.matcher.similarities', unavailable)
    partial = match_resume(candidate, job)
    assert partial.score is None and partial.components['skills'] is not None
    assert any('MiniLM' in warning for warning in partial.warnings)


def test_title_only_not_scored(candidate):
    result = match_resume(candidate, parse_job(JobInput(company='Demo', role='Data Analyst')))
    assert result.score is None


def test_seniority_and_education(candidate, jobs):
    entry = parse_job(JobInput(**jobs[0]))
    senior = parse_job(JobInput(**jobs[2]))
    assert experience_score(candidate, entry) > experience_score(candidate, senior)
    assert education_score(candidate, entry.education_required) > education_score(candidate, senior.education_required)
