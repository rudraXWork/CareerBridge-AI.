import re
from functools import lru_cache
from io import BytesIO
from zipfile import BadZipFile, ZipFile

import pymupdf
import spacy
from docx import Document

from backend.config import MAX_TEXT, MAX_UPLOAD, SPACY_MODEL
from backend.models import CandidateProfile
from backend.nlp.skills import extract_skills, suggest_roles

HEADINGS = {
    'summary': ['summary', 'profile', 'objective', 'professional summary', 'about me'],
    'skills': ['skills', 'technical skills', 'key skills', 'technologies', 'technical proficiency'],
    'education': ['education', 'academic qualifications', 'academic background'],
    'experience': ['experience', 'work experience', 'professional experience', 'employment', 'internships'],
    'projects': ['projects', 'academic projects', 'personal projects', 'key projects', 'selected projects'],
    'certifications': ['certifications', 'certificates', 'courses', 'licenses and certifications'],
    'achievements': ['achievements', 'awards', 'honors', 'accomplishments', 'achievements & certifications'],
}
EMAIL_PATTERN = r'[A-Za-z0-9.!#$%&\x27*+/=?^_`{|}~-]+@[A-Za-z0-9-]+(?:\.[A-Za-z0-9-]+)+'


@lru_cache(maxsize=1)
def load_nlp():
    try:
        return spacy.load(SPACY_MODEL)
    except OSError:
        nlp = spacy.blank('en')
        nlp.add_pipe('sentencizer')
        return nlp


def extract_text(content, filename):
    if not content or len(content) > MAX_UPLOAD:
        raise ValueError('Upload a non-empty resume under 5 MB.')
    ext = filename.rsplit('.', 1)[-1].lower()
    try:
        if ext == 'pdf':
            with pymupdf.open(stream=content, filetype='pdf') as doc:
                if doc.needs_pass:
                    raise ValueError('Password-protected PDFs are not supported.')
                if len(doc) > 30:
                    raise ValueError('Resume PDFs must have 30 pages or fewer.')
                text = '\n'.join(page.get_text(sort=True) for page in doc)
        elif ext == 'docx':
            with ZipFile(BytesIO(content)) as archive:
                if sum(info.file_size for info in archive.infolist()) > 25 * 1024 * 1024:
                    raise ValueError('The DOCX expands beyond the document size limit.')
            doc = Document(BytesIO(content))
            text = '\n'.join(p.text for p in doc.paragraphs)
            text += '\n' + '\n'.join(' | '.join(cell.text for cell in row.cells)
                                     for table in doc.tables for row in table.rows)
        elif ext == 'txt':
            text = content.decode('utf-8-sig')
        else:
            raise ValueError('Use PDF, DOCX or UTF-8 TXT.')
    except (RuntimeError, BadZipFile, UnicodeDecodeError, KeyError) as exc:
        raise ValueError('Could not read this document. Check its format.') from exc
    if not text.strip():
        raise ValueError('No readable text found. Scanned PDFs need OCR before upload.')
    if len(text) > MAX_TEXT:
        raise ValueError('The extracted resume exceeds 50,000 characters.')
    return text


def clean_text(text):
    text = text.replace('\r\n', '\n').replace('\r', '\n')
    text = text.replace('\xa0', ' ').replace('\u200b', '').replace('\ufeff', '')
    text = text.replace('\uf0b7', '•').replace('\x00', '').replace('\xad', '')
    text = re.sub(r'[ \t]+', ' ', text)
    return re.sub(r'\n{3,}', '\n\n', text).strip()


def detect_sections(text):
    lookup = {heading: key for key, headings in HEADINGS.items() for heading in headings}
    sections = {'header': []}
    current = 'header'
    for line in text.splitlines():
        heading, sep, remainder = line.strip().partition(':')
        key = lookup.get(heading.strip(' •-').casefold())
        if key:
            current = key
            sections.setdefault(key, [])
            if sep and remainder.strip():
                sections[key].append(remainder.strip())
        elif line.strip():
            sections[current].append(line.strip())
    return {key: '\n'.join(lines) for key, lines in sections.items()}


def experience_years(text):
    # Do not add overlapping jobs or assume projects are employment.
    match = re.search(r'(\d+(?:\.\d+)?)\+?\s+years?\s+(?:of\s+)?(?:professional |work |industry )?experience', text, re.I)
    return float(match[1]) if match else None


def parse_resume(text):
    text = clean_text(text)
    sections = detect_sections(text)
    nlp = load_nlp()
    header = sections.get('header', '')
    people = [ent.text for ent in nlp(header[:1500]).ents if ent.label_ == 'PERSON']
    name = ''
    first = header.split('\n')[0].strip() if header else ''
    # NER can label only part of a full name as a person.
    if re.fullmatch(r'[A-Za-z][A-Za-z .\x27-]{2,60}', first) and 2 <= len(first.split()) <= 4:
        if not extract_skills(first)[1] and not re.search(r'resume|curriculum|engineer|analyst|developer', first, re.I):
            name = first
    if not name:
        name = people[0] if people else ''
    email = re.search(EMAIL_PATTERN, text)
    relevant = '\n'.join(sections.get(key, '') for key in ('skills', 'projects', 'experience', 'summary'))
    skills, normalized = extract_skills(relevant.strip() or text)
    warnings = ['Rule-based extraction can miss details. Review and correct the profile before matching.']
    if 'ner' not in nlp.pipe_names:
        warnings.append('spaCy English model is missing; using tokenization and name rules. Run download_models.py.')
    if not any(key in sections for key in HEADINGS):
        warnings.append('No standard section headings found; skills were extracted from all text.')
    return CandidateProfile(
        name=name, email=email.group() if email else '', summary=sections.get('summary', ''),
        education=sections.get('education', '').splitlines(), skills=skills, normalized_skills=normalized,
        projects=sections.get('projects', '').splitlines(), experience=sections.get('experience', '').splitlines(),
        technologies=normalized, certifications=sections.get('certifications', '').splitlines(),
        possible_roles=suggest_roles(normalized), sections=sections,
        experience_years=experience_years(sections.get('experience', '') + '\n' + sections.get('summary', '')),
        warnings=warnings,
    )
