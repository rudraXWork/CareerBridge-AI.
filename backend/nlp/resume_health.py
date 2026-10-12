import re
from io import BytesIO

import pymupdf
from docx import Document
from docx.oxml.ns import qn

from backend.models import ResumeFileInfo
from backend.nlp.resume_parser import EMAIL_PATTERN, detect_sections

VERBS = set(('built developed designed created implemented analyzed analysed improved reduced increased '
             'automated optimized tested deployed maintained led managed delivered researched trained '
             'evaluated cleaned queried wrote integrated organized presented collaborated launched '
             'captured engineered').split())
METRIC = re.compile(r'\b\d[\d,]*(?:\.\d+)?\s*%|[$₹€]\s*\d|\b\d[\d,]*(?:\.\d+)?\s+'
                    r'(?:users?|records?|rows?|customers?|hours?|minutes?|seconds?|requests?|datasets?|'
                    r'clients?|students?|transactions?|downloads?|sales|reports?)\b', re.I)


def inspect_document(content, filename, text):
    info = ResumeFileInfo(format=filename.rsplit('.', 1)[-1].upper(), text_chars=len(text.strip()),
                          section_names=[key for key, body in detect_sections(text).items()
                                         if key != 'header' and body.strip()])
    if re.search(r'\b(?:[A-Z] ){4,}[A-Z]\b', text):
        info.formatting_issues.append('Spaced-out lettering detected; use normal heading text.')
    if re.search(r"\b(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)\.?\s+'?\d{2}\b", text, re.I):
        info.formatting_issues.append('Short date years detected; use four-digit years consistently.')
    if '\ufffd' in text:
        info.formatting_issues.append('Replacement characters detected; check the extracted text against the file.')
    if info.format == 'DOCX':
        doc = Document(BytesIO(content))
        info.tables = len(doc.tables)
        info.text_boxes = len(doc.element.xpath('.//w:txbxContent'))
        info.columns = any(int(col.get(qn('w:num'), '1')) > 1 for col in doc.element.xpath('.//w:sectPr/w:cols'))
        outside = '\n'.join(p.text for section in doc.sections
                            for area in (section.header, section.footer) for p in area.paragraphs)
        info.header_footer_contact = bool(re.search(EMAIL_PATTERN, outside))
    elif info.format == 'PDF':
        # Side-by-side text blocks are a possible column layout, not proof of a parse failure.
        info.columns = False
        with pymupdf.open(stream=content, filetype='pdf') as doc:
            for page in doc:
                blocks = [b for b in page.get_text('blocks') if b[6] == 0 and len(b[4].strip()) > 25]
                for i, left in enumerate(blocks):
                    for right in blocks[i + 1:]:
                        a, b = sorted((left, right), key=lambda block: block[0])
                        overlap = min(a[3], b[3]) - max(a[1], b[1])
                        if (a[2] + page.rect.width * .04 < b[0] and
                                overlap > .5 * min(a[3] - a[1], b[3] - b[1])):
                            info.columns = True
    else:
        info.tables = info.text_boxes = 0
        info.columns = info.header_footer_contact = False
    return info


def finding(key, title, weight, fraction, evidence, fix='', example=''):
    return {'id': key, 'title': title, 'weight': weight,
            'points': round(weight * fraction, 2) if fraction is not None else None,
            'status': 'unknown' if fraction is None else 'pass' if fraction == 1 else 'review',
            'evidence': evidence, 'fix': fix if fraction is not None and fraction < 1 else '',
            'example': example}


def total(checks):
    weight = sum(item['weight'] for item in checks if item['points'] is not None)
    return {'score': round(sum(item['points'] or 0 for item in checks) / weight * 100, 1) if weight else None,
            'coverage': weight, 'checks': checks}


def content_statements(lines):
    marker = re.compile(r'^(?:[•●▪*]\s*|[-–]\s+|\d+[.)]\s+)')
    text = [line.strip() for item in lines for line in item.splitlines() if line.strip()]
    statements = []
    wrapped = False
    pending = False
    has_bullets = any(marker.match(line) for line in text)
    for i, line in enumerate(text):
        marked = bool(marker.match(line)) or pending
        pending = False
        first = line.split()[0].strip('.,').lower()
        dates = re.search(r'\b(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)\w*\s+\d{4}\b', line, re.I)
        heading = (not marked and ('|' in line or (dates and first not in VERBS) or (
            len(line.split()) <= 12 and not re.search(r'[.!?]$', line)
            and i + 1 < len(text) and marker.match(text[i + 1]) and not wrapped)))
        if heading:
            wrapped = False
            continue
        body = marker.sub('', line)
        if not body:
            pending = marked
            wrapped = False
            continue
        if not marked and wrapped:
            statements[-1] += ' ' + body
        else:
            if not marked and ':' in body and not re.match(r'https?:', body, re.I):
                body = body.split(':', 1)[1].strip()
            elif has_bullets and not marked:
                # Role, employer and project headings are not achievement bullets.
                continue
            words = body.split()
            if marked or len(words) >= 5 or (words and words[0].strip('.,').lower() in VERBS):
                statements.append(body)
            else:
                wrapped = False
                continue
        wrapped = (marked or wrapped) and not re.search(r'[.!?]$', body)
    return statements


def analyze_resume(candidate):
    info = candidate.document
    ats = [finding('text', 'Extractable text', 25, float(info.text_chars > 0) if info else None,
                   f'{info.text_chars:,} text characters extracted; extraction completeness is not checked.' if info else 'Original file extraction not available.'),
           finding('contact', 'Contact details', 15,
                   (bool(candidate.name) + bool(re.fullmatch(EMAIL_PATTERN, candidate.email))) / 2,
                   f"Name {'present' if candidate.name else 'missing'}; email {'valid' if re.fullmatch(EMAIL_PATTERN, candidate.email) else 'missing or invalid'}.",
                   'Review your full name and email. Keep them as text in the document body.')]
    sections = info.section_names if info else [key for key in ('education', 'skills', 'projects', 'experience')
                                               if getattr(candidate, 'normalized_skills' if key == 'skills' else key)]
    if info and candidate.sections:
        # Older uploads may have counted headings with no extracted content.
        sections = [key for key in sections if candidate.sections.get(key, '').strip()]
    found = sum(('education' in sections, 'skills' in sections, 'projects' in sections or 'experience' in sections))
    ats.append(finding('sections', 'Recognized sections', 20, found / 3,
                       'Sections with extracted content: ' + ', '.join(sections) if sections else 'No sections with content recognized.',
                       'Use clear Education, Skills and Projects or Experience headings with content under each.'))
    for key, title, weight, fix in (
        ('tables', 'Tables', 10, 'Check text order in tables; prefer a simple text layout.'),
        ('columns', 'Columns', 10, 'Check the reading order or use a single column.'),
        ('text_boxes', 'Text boxes', 10, 'Move essential text out of text boxes into the document body.'),
        ('header_footer_contact', 'Header/footer email', 5, 'Keep your email in the document body too.'),
    ):
        value = getattr(info, key) if info else None
        evidence = ('Not checked for this file.' if value is None else
                    'Possible side-by-side columns detected.' if key == 'columns' and value else
                    f'{value} table(s) detected.' if key == 'tables' else
                    'Risk detected in the uploaded file.' if value else 'No risk detected by this check.')
        ats.append(finding(key, title, weight, None if value is None else .5 if key == 'tables' and value else float(not value),
                           evidence, fix))
    issues = info.formatting_issues if info else None
    ats.append(finding('formatting', 'Text formatting', 5, float(not issues) if issues is not None else None,
                       ' '.join(issues) if issues else 'No text artifacts detected.' if info else 'Original text not available.',
                       'Use normal lettering and four-digit years; check for garbled text.'))

    bullets = content_statements(candidate.projects) + content_statements(candidate.experience)
    count = len(bullets)
    actions = [b for b in bullets if b.split()[0].strip('.,').lower() in VERBS]
    metrics = [b for b in bullets if METRIC.search(b)]
    detailed = [b for b in bullets if len(b.split()) >= 8 and re.search(r'\b(using|with|to|for)\b', b, re.I)]
    readable = [b for b in bullets if 4 <= len(b.split()) <= 45]
    unique = len(set(re.sub(r'\W+', ' ', b.casefold()).strip() for b in bullets))
    ratio = lambda n: n / count if count else 0
    example = lambda good: next((b for b in bullets if b not in good), '')
    complete = sum((bool(candidate.education), bool(candidate.normalized_skills),
                    bool(candidate.projects or candidate.experience))) / 3
    content = [
        finding('actions', 'Action verbs', 20, ratio(len(actions)),
                f'{len(actions)} of {count} statements start with a recognized action verb.',
                'Start with a specific action such as built, analyzed or implemented, where accurate.', example(actions)),
        finding('metrics', 'Measurable results', 25, min(ratio(len(metrics)) / .3, 1),
                f'{len(metrics)} of {count} statements contain a measurable result or scale.',
                'Add a real result or scale, such as records analyzed or time saved. Do not invent numbers.', example(metrics)),
        finding('detail', 'Project and experience detail', 20, ratio(len(detailed)),
                f'{len(detailed)} of {count} statements give detail and how or why the work was done.',
                'Explain what you did, the tools or method, and its purpose.', example(detailed)),
        finding('readability', 'Statement length', 10, ratio(len(readable)),
                f'{len(readable)} of {count} statements are between 4 and 45 words.',
                'Use short, complete statements; split very long bullets.', example(readable)),
        finding('repetition', 'Repeated statements', 10, ratio(unique),
                f'{unique} distinct statements out of {count}.' +
                (f' Repeated: “{next(b for b in bullets if bullets.count(b) > 1)}”' if len(set(bullets)) < count else ''),
                'Remove repeated wording and give each bullet a different contribution.'),
        finding('completeness', 'Core content', 15, complete,
                'Checks education, skills, and projects or work experience. Student projects count.',
                'Add any missing education, skills and project or work evidence.'),
    ]
    ats_result, content_result = total(ats), total(content)
    scores = [item['score'] for item in (ats_result, content_result) if item['score'] is not None]
    return {'score': round(sum(scores) / len(scores), 1) if scores else None,
            'ats': ats_result, 'content': content_result,
            'method': 'local-rules-v2',
            'note': 'Local heuristic estimate, not a score from an employer’s ATS. Unknown checks are excluded.',
            'file_note': 'Profile edits recheck reviewed text. Re-upload an updated file to recheck its extraction and layout.'}
